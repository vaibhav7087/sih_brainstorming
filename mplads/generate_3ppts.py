#!/usr/bin/env python3
import sys, os, json
sys.path.insert(0, ".")
from gen_helpers import *
from pptx.util import Inches, Pt
from gen_slide12 import build_slide1, build_slide2
from gen_slide34 import build_slide3, build_slide4
from gen_slide56 import build_slide5, build_slide6, update_footers, fix_team_name_ovals, clean_slide7

TEMPLATE = "SIH2026-IDEA-Presentation-Format.pptx"
OUT = "generated_ppts"


def generate_option(opt_key, filename):
    print(f"\n{'='*60}")
    print(f"Generating: {filename}")
    print(f"{'='*60}")
    prs = Presentation(TEMPLATE)
    print(f"  Template loaded: {len(prs.slides)} slides")

    build_slide1(prs, opt_key)
    print("  Slide 1: Title - DONE")
    build_slide2(prs, opt_key)
    print("  Slide 2: Solution - DONE")
    build_slide3(prs, opt_key)
    print("  Slide 3: Technical Approach - DONE")
    build_slide4(prs, opt_key)
    print("  Slide 4: Feasibility - DONE")
    build_slide5(prs, opt_key)
    print("  Slide 5: Impact & Benefits - DONE")
    build_slide6(prs, opt_key)
    print("  Slide 6: References - DONE")
    update_footers(prs)
    print("  Footers updated - DONE")
    fix_team_name_ovals(prs)
    print("  Team name ovals fixed - DONE")
    clean_slide7(prs)
    print("  Slide 7 cleaned - DONE")

    outpath = os.path.join(OUT, filename)
    prs.save(outpath)
    print(f"  SAVED: {outpath}")

    prs2 = Presentation(outpath)
    print(f"  VERIFIED: {len(prs2.slides)} slides in output")
    for i, slide in enumerate(prs2.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t and len(t) > 3:
                        texts.append(t[:60])
        if texts:
            print(f"    Slide {i+1}: {texts[0]}")
    return outpath


if __name__ == "__main__":
    options = [
        ("OptionA", "MPLADS_AI_Sentinel_OptionA_HighTech.pptx"),
        ("OptionB", "MPLADS_AI_Sentinel_OptionB_BusinessImpact.pptx"),
        ("OptionC", "MPLADS_AI_Sentinel_OptionC_Innovation.pptx"),
    ]
    results = []
    for key, fname in options:
        path = generate_option(key, fname)
        results.append((key, fname, path))

    print(f"\n{'='*60}")
    print("ALL 3 PPT OPTIONS GENERATED SUCCESSFULLY")
    print(f"{'='*60}")
    for key, fname, path in results:
        print(f"  {key}: {path}")
