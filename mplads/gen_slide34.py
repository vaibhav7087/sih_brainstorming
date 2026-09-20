#!/usr/bin/env python3
import sys, os, json
sys.path.insert(0, ".")
from gen_helpers import *
from pptx.util import Inches, Pt

with open("ppt_content.json", "r", encoding="utf-8") as f:
    CONTENT = json.load(f)


def build_slide3(prs, opt):
    slide = prs.slides[2]
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if "Technologies to be used" in txt or "Methodology" in txt:
                shape.text_frame.clear()
    data = CONTENT["slide3"][opt]
    add_colored_rect(slide, Inches(0.3), Inches(1.1), Inches(6.0), Inches(0.35), AB)
    add_text(slide, Inches(0.4), Inches(1.12), Inches(5.8), Inches(0.3),
             data["left_title"], fs=Pt(13), bold=True, clr=WH)
    add_bullets(slide, Inches(0.3), Inches(1.5), Inches(6.0), Inches(4.0),
                data["left_items"], fs=Pt(11), clr=DT)
    add_colored_rect(slide, Inches(6.5), Inches(1.1), Inches(6.0), Inches(0.35), AG)
    add_text(slide, Inches(6.6), Inches(1.12), Inches(5.8), Inches(0.3),
             data["right_title"], fs=Pt(13), bold=True, clr=WH)
    add_bullets(slide, Inches(6.5), Inches(1.5), Inches(6.0), Inches(4.0),
                data["right_items"], fs=Pt(11), clr=DT)
    add_img(slide, shots["arch"], Inches(3.5), Inches(5.2), w=Inches(5.5))


def build_slide4(prs, opt):
    slide = prs.slides[3]
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if "Analysis of the feasibility" in txt or "Potential challenges" in txt:
                shape.text_frame.clear()
    data = CONTENT["slide4"][opt]
    add_colored_rect(slide, Inches(0.3), Inches(1.1), Inches(6.0), Inches(0.35), AO)
    add_text(slide, Inches(0.4), Inches(1.12), Inches(5.8), Inches(0.3),
             data["left_title"], fs=Pt(13), bold=True, clr=WH)
    add_bullets(slide, Inches(0.3), Inches(1.5), Inches(6.0), Inches(4.0),
                data["left_items"], fs=Pt(11), clr=DT)
    add_colored_rect(slide, Inches(6.5), Inches(1.1), Inches(6.0), Inches(0.35), AR)
    add_text(slide, Inches(6.6), Inches(1.12), Inches(5.8), Inches(0.3),
             data["right_title"], fs=Pt(13), bold=True, clr=WH)
    add_bullets(slide, Inches(6.5), Inches(1.5), Inches(6.0), Inches(4.0),
                data["right_items"], fs=Pt(11), clr=DT)


print("Part B loaded: Slides 3-4 builders")
