#!/usr/bin/env python3
import copy, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TEMPLATE = "SIH2026-IDEA-Presentation-Format.pptx"
OUT = "generated_ppts"
SS = "ppt_assets/screenshots"
os.makedirs(OUT, exist_ok=True)

DN = RGBColor(0x0B,0x1D,0x3A)
AB = RGBColor(0x1A,0x73,0xE8)
AG = RGBColor(0x0D,0x92,0x76)
AR = RGBColor(0xD9,0x3C,0x3E)
AO = RGBColor(0xE8,0x6C,0x00)
WH = RGBColor(0xFF,0xFF,0xFF)
DT = RGBColor(0x1A,0x1A,0x2E)
SG = RGBColor(0x66,0x66,0x66)
LB = RGBColor(0xE8,0xF0,0xFE)

shots = {
    "arch": os.path.join(SS, "architecture.png"),
    "ministry": os.path.join(SS, "ministry_dashboard.png"),
    "mp_dash": os.path.join(SS, "mp_dashboard.png"),
    "da_cag": os.path.join(SS, "da_cag_dashboard.png"),
    "ml_eng": os.path.join(SS, "ml_engine.png"),
    "login": os.path.join(SS, "login.png"),
    "userflow": os.path.join(SS, "userflow.png"),
}


def add_bullets(slide, left, top, w, h, items, fs=Pt(12), clr=DT):
    tb = slide.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "\u2022  " + item
        p.font.name = "Arial"
        p.font.size = fs
        p.font.color.rgb = clr
        p.space_before = Pt(2)
        p.space_after = Pt(2)


def add_img(slide, path, left, top, w=None, h=None):
    if os.path.exists(path):
        try:
            if w and h:
                slide.shapes.add_picture(path, left, top, w, h)
            elif w:
                slide.shapes.add_picture(path, left, top, width=w)
            elif h:
                slide.shapes.add_picture(path, left, top, height=h)
            else:
                slide.shapes.add_picture(path, left, top)
            return True
        except Exception:
            pass
    return False


def add_text(slide, left, top, w, h, text, fs=Pt(12), bold=False, clr=DT, align=PP_ALIGN.LEFT, fn="Arial"):
    tb = slide.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = fn
    p.font.size = fs
    p.font.bold = bold
    p.font.color.rgb = clr
    p.alignment = align


def add_multiline(slide, left, top, w, h, lines, fs=Pt(11), clr=DT, bold_first=False, spacing=Pt(4)):
    tb = slide.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Arial"
        p.font.size = fs
        p.font.color.rgb = clr
        if bold_first and i == 0:
            p.font.bold = True
        p.space_before = spacing
        p.space_after = spacing


def add_colored_rect(slide, left, top, w, h, fill_clr):
    shape = slide.shapes.add_shape(
        1, left, top, w, h  # MSO_SHAPE.RECTANGLE = 1
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_clr
    shape.line.fill.background()
    return shape


def add_rounded_rect(slide, left, top, w, h, fill_clr):
    shape = slide.shapes.add_shape(
        5, left, top, w, h  # MSO_SHAPE.ROUNDED_RECTANGLE = 5
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_clr
    shape.line.fill.background()
    return shape


def add_card(slide, left, top, w, h, title, items, title_clr=AB, bg_clr=LB):
    card = add_rounded_rect(slide, left, top, w, h, bg_clr)
    add_text(slide, left + Inches(0.15), top + Inches(0.08), w - Inches(0.3), Inches(0.35),
             title, fs=Pt(11), bold=True, clr=title_clr)
    add_bullets(slide, left + Inches(0.15), top + Inches(0.4), w - Inches(0.3), h - Inches(0.5),
                items, fs=Pt(9), clr=DT)


print("Helper functions loaded successfully.")
