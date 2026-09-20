#!/usr/bin/env python3
import sys, os, json
sys.path.insert(0, ".")
from gen_helpers import *
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

TEAM = "MPLADS AI Sentinel"
PS_ID = "26102"
PS_TITLE = "Smart Automation of MPLADS Fund Monitoring using Explainable AI"
THEME = "Smart Automation"
CAT = "Software"

with open("ppt_content.json", "r", encoding="utf-8") as f:
    CONTENT = json.load(f)


def build_slide1(prs, opt):
    slide = prs.slides[0]
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if "TITLE PAGE" in txt:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = TEAM
                p.font.name = "Times New Roman"
                p.font.size = Pt(28)
                p.font.bold = True
                p.font.color.rgb = DN
                p.alignment = PP_ALIGN.LEFT
            elif "SMART INDIA HACKATHON 2026" in txt:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = "SMART INDIA HACKATHON 2026"
                p.font.name = "Garamond"
                p.font.size = Pt(40)
                p.font.bold = True
                p.font.color.rgb = DN
            elif "Problem Statement ID" in txt:
                shape.text_frame.clear()
                lines = [
                    "Problem Statement ID: " + PS_ID,
                    "Problem Statement Title: " + PS_TITLE,
                    "Theme: " + THEME,
                    "PS Category: " + CAT,
                    "Team ID: " + TEAM,
                    "Team Name: " + TEAM,
                ]
                for i, line in enumerate(lines):
                    p = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
                    p.text = line
                    p.font.name = "Arial"
                    p.font.size = Pt(16)
                    p.font.bold = True
                    p.font.color.rgb = DT
                    p.space_before = Pt(4)
                    p.space_after = Pt(4)


def build_slide2(prs, opt):
    slide = prs.slides[1]
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if "IDEA TITLE" in txt:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = "MPLADS AI Sentinel: Explainable AI for Parliamentary Fund Oversight"
                p.font.name = "Times New Roman"
                p.font.size = Pt(28)
                p.font.bold = True
                p.font.color.rgb = DN
            elif "Proposed Solution" in txt or "Detailed explanation" in txt:
                shape.text_frame.clear()
    data = CONTENT["slide2"][opt]
    add_text(slide, Inches(0.5), Inches(1.2), Inches(12), Inches(0.4),
             data["subtitle"], fs=Pt(18), bold=True, clr=AB)
    add_bullets(slide, Inches(0.5), Inches(1.7), Inches(12), Inches(4.5),
                data["bullets"], fs=Pt(13), clr=DT)
    add_img(slide, shots["arch"], Inches(8.5), Inches(1.5), w=Inches(4.5))


print("Part A loaded: Slides 1-2 builders")
