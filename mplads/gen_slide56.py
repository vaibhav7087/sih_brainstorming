#!/usr/bin/env python3
import sys, os, json
sys.path.insert(0, ".")
from gen_helpers import *
from pptx.util import Inches, Pt

with open("ppt_content.json", "r", encoding="utf-8") as f:
    CONTENT = json.load(f)


def build_slide5(prs, opt):
    slide = prs.slides[4]
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if "Potential impact" in txt or "Benefits of the solution" in txt:
                shape.text_frame.clear()
    data = CONTENT["slide5"][opt]
    add_colored_rect(slide, Inches(0.3), Inches(1.1), Inches(6.0), Inches(0.35), AG)
    add_text(slide, Inches(0.4), Inches(1.12), Inches(5.8), Inches(0.3),
             data["left_title"], fs=Pt(13), bold=True, clr=WH)
    add_bullets(slide, Inches(0.3), Inches(1.5), Inches(6.0), Inches(4.0),
                data["left_items"], fs=Pt(11), clr=DT)
    add_colored_rect(slide, Inches(6.5), Inches(1.1), Inches(6.0), Inches(0.35), AB)
    add_text(slide, Inches(6.6), Inches(1.12), Inches(5.8), Inches(0.3),
             data["right_title"], fs=Pt(13), bold=True, clr=WH)
    add_bullets(slide, Inches(6.5), Inches(1.5), Inches(6.0), Inches(4.0),
                data["right_items"], fs=Pt(11), clr=DT)
    add_img(slide, shots["ministry"], Inches(2.5), Inches(5.2), w=Inches(4.0))
    add_img(slide, shots["mp_dash"], Inches(7.0), Inches(5.2), w=Inches(4.0))


def build_slide6(prs, opt):
    slide = prs.slides[5]
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if "Details / Links" in txt:
                shape.text_frame.clear()
    data = CONTENT["slide6"][opt]
    add_bullets(slide, Inches(0.5), Inches(1.3), Inches(12), Inches(4.5),
                data["items"], fs=Pt(11), clr=DT)


def update_footers(prs):
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if "@SIH Idea submission" in txt:
                    shape.text_frame.clear()
                    p = shape.text_frame.paragraphs[0]
                    p.text = "MPLADS AI Sentinel - SIH 2026"
                    p.font.name = "Arial"
                    p.font.size = Pt(10)
                    p.font.color.rgb = SG


def fix_team_name_ovals(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if "Your Team Name" in para.text:
                        for run in para.runs:
                            if "Your Team Name" in run.text:
                                run.text = run.text.replace("Your Team Name", "MPLADS AI Sentinel")


def clean_slide7(prs):
    if len(prs.slides) > 6:
        slide = prs.slides[6]
        for shape in list(slide.shapes):
            if shape.has_text_frame:
                shape.text_frame.clear()


print("Part C loaded: Slides 5-6 builders")
